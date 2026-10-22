"""The downloadable Excel template of the cash count (ops/plantillas.py): what is published is what the script builds, and its
formulas add up exactly like the calculator on the website (the same example is checked in tests/web-herramientas.test.mjs)."""

import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "ops"))

import plantillas  # noqa: E402

PUBLISHED = ROOT / "docs" / "descargas" / "plantilla-arqueo-de-caja.xlsx"


def evaluate(ws, address, memo=None):
    """A tiny evaluator for the formulas the template uses (+ - * SUM IF ROUND and comparisons): enough to check the sums
    without a spreadsheet program."""
    memo = {} if memo is None else memo
    if address in memo:
        return memo[address]
    value = ws[address].value
    if not (isinstance(value, str) and value.startswith("=")):
        memo[address] = 0 if value is None else value
        return memo[address]

    def token(match):
        text = match.group(0)
        if text.startswith('"'):
            return text
        if ":" in text:
            first, last = text.split(":")
            return f"rng({first!r},{last!r})"
        return f"ref({text!r})"

    expr = re.sub(r'"[^"]*"|[A-Z]{1,2}\d+:[A-Z]{1,2}\d+|[A-Z]{1,2}\d+', token, value[1:])
    expr = re.sub(r"(?<![<>!=])=(?!=)", "==", expr).replace("<>", "!=")

    def rng(first, last):
        cells = ws[f"{first}:{last}"]
        return [evaluate(ws, c.coordinate, memo) for row in cells for c in row]

    scope = {"ref": lambda a: evaluate(ws, a, memo), "rng": rng, "SUM": sum, "ROUND": lambda x, n: round(x, n),
             "IF": lambda condition, yes, no: yes if condition else no}
    memo[address] = eval(expr, {"__builtins__": {}}, scope)  # noqa: S307 (only the template's own formulas)
    return memo[address]


def filled(**amounts):
    wb = load_workbook(PUBLISHED)
    ws = wb[plantillas.SHEET]
    # 3 × 50 €, 7 × 20 €, 4 × 10 €, 6 × 5 €, 9 × 2 €, 11 × 1 €, 8 × 0,50 €, 12 × 0,10 € = 394,20 €
    for address, quantity in {"C10": 3, "C11": 7, "C12": 4, "C13": 6, "C14": 9, "C15": 11, "C16": 8, "C18": 12}.items():
        ws[address] = quantity
    for address, amount in amounts.items():
        ws[address] = amount
    return ws


def test_the_published_template_is_what_the_script_builds():
    # If this fails: run «python ops/plantillas.py» and commit the result.
    assert plantillas.main(["--check"]) == 0


def test_the_template_has_every_euro_banknote_and_coin_with_its_value():
    ws = load_workbook(PUBLISHED)[plantillas.SHEET]
    values = [ws[f"B{row}"].value for row in range(7, 22)]
    assert values == [500, 200, 100, 50, 20, 10, 5, 2, 1, 0.5, 0.2, 0.1, 0.05, 0.02, 0.01]
    assert [ws[f"A{row}"].value for row in (7, 14, 21)] == ["Billete de 500 €", "Moneda de 2 €", "Moneda de 0,01 €"]


def test_the_template_adds_up_the_count_the_expected_cash_and_the_difference():
    ws = filled(D25=150, D26=250, D33=100)
    assert round(evaluate(ws, "D22"), 2) == 394.20          # counted
    assert round(evaluate(ws, "D29"), 2) == 400.00          # expected = 150 + 250
    assert round(evaluate(ws, "D32"), 2) == -5.80           # counted − expected
    assert evaluate(ws, "E32") == "Falta dinero"
    assert round(evaluate(ws, "D34"), 2) == 294.20          # counted − what stays for tomorrow


def test_the_template_says_the_cash_balances_when_it_does_and_when_it_does_not():
    assert evaluate(filled(D25=144.20, D26=250), "E32") == "La caja cuadra"
    assert evaluate(filled(D25=100, D26=250), "E32") == "Sobra dinero"
    # inflows add and outflows subtract
    ws = filled(D25=100, D26=250, D27=30, D28=35.80)
    assert round(evaluate(ws, "D29"), 2) == 344.20
    assert round(evaluate(ws, "D32"), 2) == 50.00
    # coins of cents do not leave a floating point difference
    ws = load_workbook(PUBLISHED)[plantillas.SHEET]
    ws["C18"], ws["C17"], ws["D26"] = 1, 1, 0.30            # 0,10 + 0,20 against 0,30 of sales
    assert evaluate(ws, "E32") == "La caja cuadra"


def test_the_inputs_are_marked_and_the_sheet_prints_on_one_page():
    ws = load_workbook(PUBLISHED)[plantillas.SHEET]
    assert str(ws.page_setup.paperSize) == ws.PAPERSIZE_A4 and ws.sheet_properties.pageSetUpPr.fitToPage
    assert all(ws[f"C{row}"].fill.fgColor.rgb.endswith(plantillas.INPUT) for row in range(7, 22))
    assert len(ws.data_validations.dataValidation) == 2
