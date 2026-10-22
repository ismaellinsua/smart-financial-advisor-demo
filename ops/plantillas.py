"""Plantillas descargables de la web (docs/descargas/). Hoy, una: la hoja de Excel del arqueo de caja.

    python ops/plantillas.py            # escribe docs/descargas/plantilla-arqueo-de-caja.xlsx
    python ops/plantillas.py --check    # no escribe nada: sale con error si el fichero publicado no es el que sale de aquí

La plantilla hace las mismas cuentas que la calculadora de la web (herramientas/calculadora-de-arqueo-de-caja/): lo contado por
denominaciones, lo esperado (fondo + ventas en efectivo + entradas − salidas), la diferencia y el efectivo que se retira. Es una
sola hoja, sin macros ni conexiones, y se imprime en una página A4. Los bytes del .xlsx cambian cada vez que se guarda (llevan
fechas), así que --check compara lo que importa: textos, fórmulas y formatos de cada celda.
"""

from __future__ import annotations

import argparse
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "descargas" / "plantilla-arqueo-de-caja.xlsx"
SHEET = "Arqueo de caja"
EURO = '#,##0.00" €"'

DENOMINACIONES = [(500, "Billete de 500 €"), (200, "Billete de 200 €"), (100, "Billete de 100 €"), (50, "Billete de 50 €"), (20, "Billete de 20 €"),
                  (10, "Billete de 10 €"), (5, "Billete de 5 €"), (2, "Moneda de 2 €"), (1, "Moneda de 1 €"), (0.5, "Moneda de 0,50 €"),
                  (0.2, "Moneda de 0,20 €"), (0.1, "Moneda de 0,10 €"), (0.05, "Moneda de 0,05 €"), (0.02, "Moneda de 0,02 €"), (0.01, "Moneda de 0,01 €")]

BLUE, LIGHT, INPUT, GREY = "3B5BFD", "EEF2FF", "FFF7D6", "5B6785"
thin = Side(style="thin", color="C9D3EA")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def build() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    wb.properties.creator = "NirKanA"
    wb.properties.title = "Plantilla de arqueo de caja"
    wb.properties.created = wb.properties.modified = datetime(2026, 10, 22)
    for col, width in zip("ABCDE", (50, 14, 12, 18, 22)):
        ws.column_dimensions[col].width = width

    ws["A1"] = "Arqueo de caja"
    ws["A1"].font = Font(name="Calibri", size=18, bold=True, color="0B1B33")
    ws["A2"] = "Rellena solo las celdas amarillas. Plantilla gratuita de NirKanA · nirkana.es/herramientas/calculadora-de-arqueo-de-caja/"
    ws["A2"].font = Font(name="Calibri", size=9, color=GREY)

    def label(cell, text, bold=False, color="0B1B33"):
        ws[cell] = text
        ws[cell].font = Font(name="Calibri", size=11, bold=bold, color=color)

    def entrada(cell, fmt=None, validar=None):
        ws[cell].fill = PatternFill("solid", fgColor=INPUT)
        ws[cell].border = BOX
        ws[cell].alignment = Alignment(horizontal="right")
        if fmt:
            ws[cell].number_format = fmt
        if validar:
            validar.add(cell)

    def cabecera(fila, textos):
        for col, texto in zip("ABCD", textos):
            c = ws[f"{col}{fila}"]
            c.value = texto
            c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=BLUE)
            c.alignment = Alignment(horizontal="left" if col == "A" else "right")

    enteros = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True,
                             errorTitle="Cantidad", error="Escribe un número entero de 0 o más.")
    importes = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True,
                              errorTitle="Importe", error="Escribe un importe de 0 € o más.")
    ws.add_data_validation(enteros)
    ws.add_data_validation(importes)

    label("A4", "Fecha", bold=True)
    entrada("B4")
    ws["B4"].number_format = "dd/mm/yyyy"
    label("C4", "Turno o responsable", bold=True)
    ws.merge_cells("C4:D4")
    ws["C4"].alignment = Alignment(horizontal="right")
    entrada("E4")

    cabecera(6, ("1. Efectivo contado", "Valor", "Cantidad", "Importe"))
    fila = 7
    for valor, nombre in DENOMINACIONES:
        label(f"A{fila}", nombre)
        ws[f"B{fila}"] = valor
        ws[f"B{fila}"].number_format = EURO
        ws[f"B{fila}"].font = Font(name="Calibri", size=11, color=GREY)
        entrada(f"C{fila}", "0", enteros)
        ws[f"D{fila}"] = f"=B{fila}*C{fila}"
        ws[f"D{fila}"].number_format = EURO
        ws[f"D{fila}"].border = BOX
        fila += 1
    ultima = fila - 1                                                   # 21
    label(f"A{fila}", "Total contado", bold=True)
    ws[f"D{fila}"] = f"=SUM(D7:D{ultima})"
    ws[f"D{fila}"].number_format = EURO
    ws[f"D{fila}"].font = Font(name="Calibri", size=12, bold=True)
    ws[f"D{fila}"].fill = PatternFill("solid", fgColor=LIGHT)
    contado = fila                                                      # 22

    fila += 2                                                           # 24
    cabecera(fila, ("2. Lo que debería haber", "", "", "Importe"))
    esperado_filas = {}
    for nombre, clave in (("Fondo de caja con el que empezaste", "fondo"), ("+ Ventas cobradas en efectivo", "ventas"),
                          ("+ Otras entradas de efectivo", "entradas"), ("− Salidas de efectivo (proveedores, retiradas, devoluciones)", "salidas")):
        fila += 1
        label(f"A{fila}", nombre)
        entrada(f"D{fila}", EURO, importes)
        esperado_filas[clave] = fila
    fila += 1
    label(f"A{fila}", "Efectivo esperado", bold=True)
    f = esperado_filas
    ws[f"D{fila}"] = f"=D{f['fondo']}+D{f['ventas']}+D{f['entradas']}-D{f['salidas']}"
    ws[f"D{fila}"].number_format = EURO
    ws[f"D{fila}"].font = Font(name="Calibri", size=12, bold=True)
    ws[f"D{fila}"].fill = PatternFill("solid", fgColor=LIGHT)
    esperado = fila

    fila += 2
    cabecera(fila, ("3. Resultado", "", "", "Importe"))
    fila += 1
    label(f"A{fila}", "Diferencia (contado − esperado)", bold=True)
    ws[f"D{fila}"] = f"=D{contado}-D{esperado}"
    ws[f"D{fila}"].number_format = '+#,##0.00" €";-#,##0.00" €";0.00" €"'
    ws[f"D{fila}"].font = Font(name="Calibri", size=12, bold=True)
    ws[f"E{fila}"] = f'=IF(ROUND(D{fila},2)=0,"La caja cuadra",IF(D{fila}>0,"Sobra dinero","Falta dinero"))'
    ws[f"E{fila}"].font = Font(name="Calibri", size=11, bold=True)
    diferencia = fila
    verde, ambar = PatternFill("solid", bgColor="CDEFD8"), PatternFill("solid", bgColor="FFE3B3")
    ws.conditional_formatting.add(f"D{diferencia}:E{diferencia}", FormulaRule(formula=[f"ROUND($D${diferencia},2)=0"], fill=verde))
    ws.conditional_formatting.add(f"D{diferencia}:E{diferencia}", FormulaRule(formula=[f"ROUND($D${diferencia},2)<>0"], fill=ambar))
    fila += 1
    label(f"A{fila}", "Fondo que dejas en la caja para mañana")
    entrada(f"D{fila}", EURO, importes)
    dejar = fila
    fila += 1
    label(f"A{fila}", "Efectivo a retirar del cajón", bold=True)
    ws[f"D{fila}"] = f"=D{contado}-D{dejar}"
    ws[f"D{fila}"].number_format = EURO
    ws[f"D{fila}"].font = Font(name="Calibri", size=12, bold=True)
    ws[f"D{fila}"].fill = PatternFill("solid", fgColor=LIGHT)

    fila += 2
    label(f"A{fila}", "Observaciones (qué ha pasado si no cuadra, a quién se avisa)", bold=True)
    fila += 1
    ws.merge_cells(f"A{fila}:E{fila + 2}")
    ws[f"A{fila}"].fill = PatternFill("solid", fgColor=INPUT)
    ws[f"A{fila}"].alignment = Alignment(vertical="top", wrap_text=True)
    fila += 4
    label(f"A{fila}", "Firma de quien cuenta:")
    label(f"C{fila}", "Firma de quien revisa:")
    fila += 2
    ws[f"A{fila}"] = "Las ventas con tarjeta, Bizum u otras formas de pago no están en el cajón: compara su total con el del datáfono y con el de tu programa."
    ws[f"A{fila}"].font = Font(name="Calibri", size=9, italic=True, color=GREY)

    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:E{fila}"
    ws.sheet_view.showGridLines = False
    return wb


def structure(wb: Workbook) -> dict[str, object]:
    """Lo que hace que dos plantillas sean la misma: valor o fórmula y formato de cada celda con contenido."""
    ws = wb[SHEET]
    cells = {c.coordinate: (c.value, c.number_format) for row in ws.iter_rows() for c in row if c.value is not None}
    return {"cells": cells, "merged": sorted(str(r) for r in ws.merged_cells.ranges), "print": ws.print_area}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--check", action="store_true", help="no escribe: falla si lo publicado no coincide")
    args = parser.parse_args(argv)
    wb = build()
    if args.check:
        if not OUT.exists():
            print(f"falta {OUT.relative_to(ROOT)}", file=sys.stderr)
            return 1
        if structure(load_workbook(OUT)) != structure(load_workbook_from(wb)):
            print(f"desactualizado: {OUT.relative_to(ROOT)}", file=sys.stderr)
            return 1
        return 0
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"{OUT.relative_to(ROOT)}: {OUT.stat().st_size // 1024 or 1} KB")
    return 0


def load_workbook_from(wb: Workbook) -> Workbook:
    """Pasa la hoja por un guardado en memoria, igual que el fichero publicado, para comparar lo mismo con lo mismo."""
    import io
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return load_workbook(buffer)


if __name__ == "__main__":
    sys.exit(main())
